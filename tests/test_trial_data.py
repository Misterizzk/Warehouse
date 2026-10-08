"""Trial dataset guarantees: preserve source, cover scenarios, and import exactly."""
from datetime import date, timedelta
import importlib.util
import json
from pathlib import Path
import shutil
import sqlite3
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from warehouse.database import connect, seed_demo
from warehouse.dashboard import overview
from warehouse.domain import aging, location_code
from warehouse.inventory_excel import FIELDS, HEADERS, read_inventory
from warehouse.sessions import Sessions
from warehouse.recommendation_data import recommendation_context

SCRIPT = Path(__file__).resolve().parents[1]/'scripts'/'create_trial_data.py'
SPEC = importlib.util.spec_from_file_location('create_trial_data', SCRIPT)
trial = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(trial)


class TrialDataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory(prefix='warehouse-trial-test-')
        cls.root = Path(cls.folder.name)
        cls.today = date.today()
        source = connect(cls.root/'data'/'warehouse_demo.db')
        seed_demo(source, today=cls.today-timedelta(days=12))
        with source:
            for column in range(1, 4):
                for row in range(1, 4):
                    for level in range(1, 4):
                        source.execute('INSERT INTO locations VALUES(?,?,?,?,?,?)',
                                       (location_code('C','C01',column,row,level),'C','C01',column,row,level))
            source.execute('INSERT INTO site_demand VALUES(?,?,?,?,?,?)',
                           ('FLT-001','A',1,1,0,99))
            source.execute('INSERT INTO material_unit_compatibility VALUES(?,?)',
                           ('FLT-001','UNIT-CUSTOMER-B'))
            source.execute('INSERT INTO customer_demand VALUES(?,?,?)',
                           ('B','UNIT-CUSTOMER-B',0))
        sessions = Sessions(source)
        key = sessions.start('Petugas uji', 'A')
        sessions.add_code(key, 'ID:DEMO-0001')
        sessions.save_note(key, 'Catatan sumber harus tetap sama.')
        sessions.finish(key)
        sessions.start('Petugas draft', 'B')
        cls.originals = {table: trial.table_rows(source, table) for table in (
            'batches','locations','metadata','check_sessions','check_results','stock_movements',
            'batch_location_migration','site_demand','material_unit_compatibility','customer_demand')}
        source.close()

        book = Workbook()
        ws = book.active
        ws.title = 'Inventory'
        ws.append(['Plant','Storage Location','Material','Material Description','GR Date','Storage Bin','Total Stock','Aging','Status'])
        for i in range(1, 7):
            ws.append(['BHT','PBHT',f'0000{i:03d}',f'LEGACY PART {i}',
                       cls.today-timedelta(days=10*i),f'00007860{i:02d}',i,None,None])
        book.save(cls.root/'inventory_dummy.xlsx')
        book.close()
        legacy = sqlite3.connect(cls.root/'warehouse.db')
        legacy.executescript('''CREATE TABLE sessions(id INTEGER PRIMARY KEY, officer TEXT, started TEXT, state TEXT);
            CREATE TABLE scans(id INTEGER PRIMARY KEY,session_id INTEGER,item_key TEXT,snapshot TEXT,age INTEGER,status TEXT);
            INSERT INTO sessions VALUES(1,'Petugas lama','2026-01-01T10:00:00','Draft');
            INSERT INTO scans VALUES(1,1,'00001','{"Material":"00001"}',100,'Melewati batas');''')
        legacy.commit()
        legacy.close()
        cls.hashes = {str(path.relative_to(cls.root)): trial.file_hash(path) for path in (
            cls.root/'data'/'warehouse_demo.db',cls.root/'inventory_dummy.xlsx',cls.root/'warehouse.db')}
        cls.manifest = trial.build_trial_dataset(root=cls.root,today=cls.today)

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def seed(self):
        db = sqlite3.connect(self.root/'data_uji'/'warehouse_demo_seed.db')
        db.row_factory = sqlite3.Row
        self.addCleanup(db.close)
        return db

    def test_sources_and_historical_values_are_unchanged(self):
        for path, expected in self.hashes.items():
            self.assertEqual(trial.file_hash(self.root/path), expected)
        seed = self.seed()
        for table, original in self.originals.items():
            new = trial.table_rows(seed,table)
            for record in original:
                self.assertIn(record,new,table)
        book = load_workbook(self.root/'data_uji'/'Data_Uji_Warehouse.xlsx',read_only=True)
        self.addCleanup(book.close)
        self.assertEqual(book['Barang_Asli'].max_row,len(self.originals['batches'])+1)
        self.assertEqual(book['Hasil'].cell(2,4).value,self.originals['check_results'][0]['snapshot'])
        self.assertEqual(book['Hasil_Lama'].cell(2,4).value,'{"Material":"00001"}')

    def test_scenarios_have_expected_dashboard_semantics(self):
        db = self.seed()
        summary = overview(db,today=self.today)
        self.assertEqual(db.execute('SELECT COUNT(*) FROM batches').fetchone()[0],200)
        self.assertEqual(db.execute('SELECT COUNT(*) FROM locations').fetchone()[0],324)
        self.assertEqual(len(summary['racks']),12)
        for rack,status in [('C02','Normal'),('C03','Warning'),('C04','Melewati batas')]:
            self.assertEqual(summary['racks'][('C',rack)]['status'],status)
            self.assertEqual({r['status'] for r in summary['racks'][('C',rack)]['rows']},{status})
        self.assertEqual(summary['racks'][('D','D01')]['rows'],[])
        zero=summary['racks'][('D','D02')]
        self.assertGreater(len(zero['rows']),0)
        self.assertEqual((zero['status'],zero['batches'],zero['stock']),('Tanpa stok',0,0))
        for days in (0,29,30,31,89,90,91):
            row=db.execute('SELECT * FROM batches WHERE id=?',(f'UJI-AGE-{days:03d}',)).fetchone()
            self.assertEqual(aging(row['received_date'],today=self.today)[0],days)
        for level,status in [(1,'Normal'),(2,'Warning'),(3,'Melewati batas')]:
            cell=[c for c in summary['racks'][('C','C01')]['sections'][(1,1)]['cells'] if c['level']==level][0]
            self.assertEqual(cell['status'],status)

    def test_pic_recommendation_scenarios_have_named_site_and_customer_evidence(self):
        # Each PIC example must be demonstrable from stored facts, not random names.
        db = self.seed()
        cases = [
            ('UJI-REC-001', 'Normal', 'A', None),
            ('UJI-REC-002', 'Normal', None, None),
            ('UJI-REC-003', 'Warning', None, 'A'),
            ('UJI-REC-004', 'Melewati batas', None, 'A'),
            ('UJI-REC-005', 'Melewati batas', 'B', None),
            ('UJI-REC-006', 'Melewati batas', 'C', 'A'),
        ]
        for batch_id, status, site, customer in cases:
            row = db.execute('SELECT material,received_date FROM batches WHERE id=?',
                             (batch_id,)).fetchone()
            self.assertIsNotNone(row, batch_id)
            self.assertEqual(aging(row['received_date'], today=self.today)[1], status)
            context = recommendation_context(db, row['material'])
            self.assertEqual(context['source'], 'Simulasi')
            candidate_sites = [record['site'] for record in context['sites']
                               if record['open_demand'] > record['site_stock']
                               and record['demand_now'] - record['demand_prev'] >= 2]
            candidate_customers = [record['customer'] for record in context['customers']
                                   if record['open_demand'] > 0]
            self.assertEqual(site in candidate_sites if site else not candidate_sites, True)
            self.assertEqual(customer in candidate_customers if customer else not candidate_customers,
                             True)
        self.assertEqual(recommendation_context(db, 'MATERIAL-TIDAK-DIKENAL'),
                         {'sites': [], 'customers': [], 'source': 'Simulasi'})

    def test_every_active_trial_material_has_explicit_demo_destinations(self):
        # Random clicks on existing demo inventory should show evidence, while
        # a newly imported material still receives no invented destination.
        db = self.seed()
        statuses_by_material = {}
        for row in db.execute('SELECT material,received_date FROM batches WHERE stock>0'):
            statuses_by_material.setdefault(row['material'], set()).add(
                aging(row['received_date'], today=self.today)[1])
        for material, statuses in statuses_by_material.items():
            context = recommendation_context(db, material)
            self.assertEqual({record['site'] for record in context['sites']},
                             {'A','B','C'}, material)
            if (statuses & {'Warning','Melewati batas'} and material != 'MAT-REC-005'):
                self.assertTrue(any(record['open_demand']>0 for record in context['customers']),
                                material)

    def test_workbook_roundtrip_preserves_codes_and_has_no_formulas(self):
        path=self.root/'data_uji'/'Data_Uji_Warehouse.xlsx'
        records=read_inventory(path)
        self.assertEqual(len(records),200)
        self.assertFalse(any(r['issues'] for r in records))
        coded=[r['values'] for r in records if r['values']['id']=='UJI-TEXT-001'][0]
        self.assertEqual((coded['material'],coded['plant'],coded['storage_location']),('0000123405','0010','0001'))
        self.assertEqual(self.manifest['validation'],{'imported_batches':200,'duplicate_ids_rejected':200})
        book=load_workbook(path,read_only=True)
        self.addCleanup(book.close)
        self.assertEqual(tuple(c.value for c in book['Inventory'][1]),HEADERS)
        for sheet in book:
            for row in sheet:
                self.assertFalse(any(c.data_type in ('f','e') for c in row),sheet.title)
        for row in book['Pemetaan_Bin'].iter_rows(min_row=2):
            self.assertTrue(row[2].value.startswith('00007860'))
            self.assertIn('SIMULASI',row[5].value)
        self.assertIn('Demand_Site_Simulasi', book.sheetnames)
        self.assertIn('Customer_Simulasi', book.sheetnames)
        site_rows = list(book['Demand_Site_Simulasi'].values)
        self.assertIn(('MAT-REC-001','A',3,9,10,1), site_rows)
        customer_rows = list(book['Customer_Simulasi'].values)
        self.assertIn(('A','UNIT-CUSTOMER-A-WARNING',4), customer_rows)

    def test_rerun_from_archived_legacy_files_matches_the_same_records(self):
        with tempfile.TemporaryDirectory(prefix='warehouse-trial-archive-') as folder:
            root=Path(folder)
            (root/'data').mkdir()
            archive=root/'arsip'/'pengembangan_lama'
            archive.mkdir(parents=True)
            shutil.copy2(self.root/'data'/'warehouse_demo.db',root/'data'/'warehouse_demo.db')
            for name in ('inventory_dummy.xlsx','warehouse.db'):
                shutil.copy2(self.root/name,archive/name)
            manifest=trial.build_trial_dataset(root=root,today=self.today)
            self.assertEqual(manifest['inventory_batches'],200)
            new=sqlite3.connect(root/'data_uji'/'warehouse_demo_seed.db')
            new.row_factory=sqlite3.Row
            try:
                self.assertEqual(trial.table_rows(new,'batches'),trial.table_rows(self.seed(),'batches'))
            finally:
                new.close()
            self.assertIn(str(Path('arsip/pengembangan_lama/inventory_dummy.xlsx')),manifest['source_sha256'])

    def test_legacy_exact_match_does_not_duplicate_an_existing_batch(self):
        with tempfile.TemporaryDirectory(prefix='warehouse-trial-match-') as folder:
            root=Path(folder)
            (root/'data').mkdir()
            for name in ('inventory_dummy.xlsx','warehouse.db'):
                shutil.copy2(self.root/name,root/name)
            shutil.copy2(self.root/'data'/'warehouse_demo.db',root/'data'/'warehouse_demo.db')
            original=self.originals['batches'][0]
            book=load_workbook(root/'inventory_dummy.xlsx')
            values=[original['plant'],original['storage_location'],original['material'],original['description'],
                    original['received_date'],'0000001234',original['stock'],None,None]
            for col,value in enumerate(values,1):
                book['Inventory'].cell(2,col,value)
            book.save(root/'inventory_dummy.xlsx')
            book.close()
            manifest=trial.build_trial_dataset(root=root,today=self.today)
            self.assertEqual(manifest['legacy_batches_added'],5)
            wb=load_workbook(root/'data_uji'/'Data_Uji_Warehouse.xlsx',read_only=True)
            try:
                self.assertEqual(wb['Pemetaan_Bin']['D2'].value,original['id'])
                self.assertEqual(wb['Pemetaan_Bin']['C2'].value,'0000001234')
            finally:
                wb.close()


if __name__=='__main__':
    unittest.main()
