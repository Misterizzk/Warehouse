"""Inspection reports retain the recommendation shown when an item was checked."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
import pypdfium2 as pdfium

from warehouse.database import connect, seed_demo
from warehouse.recommendations import RecommendationService
from warehouse.reports import export_excel, export_pdf, report_data
from warehouse.sessions import Sessions


class RecommendationHistoryTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.db = connect(Path(self.folder.name) / "history.db")
        seed_demo(self.db, today=date.today())

    def tearDown(self):
        self.db.close()
        self.folder.cleanup()

    def test_finished_report_keeps_site_advice_when_demo_demand_changes(self):
        with self.db:
            self.db.execute(
                """INSERT INTO site_demand
                   (material, site, demand_prev, demand_now, open_demand, site_stock)
                   VALUES ('MAT-004', 'A', 2, 8, 9, 1)"""
            )
        sessions = Sessions(self.db)
        key = sessions.start("Petugas Uji", "Area A")
        sessions.add_code(key, "DEMO-0004")
        sessions.finish(key)

        original = report_data(self.db, key)["rows"][0]["recommendation"]
        self.assertIn("Site A", original["next_action"])

        with self.db:
            self.db.execute(
                "UPDATE site_demand SET demand_now=0, open_demand=0 WHERE material='MAT-004' AND site='A'"
            )
        historical = report_data(self.db, key)["rows"][0]["recommendation"]
        self.assertEqual(historical, original)

    def test_report_exports_the_destination_selected_by_officer(self):
        with self.db:
            self.db.execute(
                """INSERT INTO site_demand
                   (material, site, demand_prev, demand_now, open_demand, site_stock)
                   VALUES ('MAT-004', 'A', 2, 8, 9, 1)"""
            )
        sessions = Sessions(self.db)
        key = sessions.start("Petugas Uji", "Area A")
        sessions.add_code(key, "DEMO-0004")
        RecommendationService(self.db).save_decision(
            key, "DEMO-0004", "Layak", "Pindahkan ke site", "Site A membutuhkan barang.",
            target_kind="site", target_code="A"
        )
        sessions.finish(key)

        data = report_data(self.db, key)
        self.assertEqual(data["rows"][0]["decision"]["target_kind"], "site")
        self.assertEqual(data["rows"][0]["decision"]["target_code"], "A")
        output = Path(self.folder.name) / "report.xlsx"
        export_excel(data, output)
        book = load_workbook(output, read_only=True, data_only=True)
        try:
            header, item = list(book["Hasil pengecekan"].values)[:2]
            self.assertEqual(dict(zip(header, item))["Tujuan tindakan"], "Site A")
        finally:
            book.close()

    def test_pdf_keeps_green_site_recommendation_even_without_staff_decision(self):
        with self.db:
            self.db.execute('INSERT INTO site_demand VALUES(?,?,?,?,?,?)',
                            ('MAT-004', 'A', 2, 8, 9, 1))
        sessions = Sessions(self.db)
        key = sessions.start('Petugas Uji', 'Area A')
        sessions.add_code(key, 'DEMO-0004')
        sessions.finish(key)
        data = report_data(self.db, key)
        self.assertEqual(data['rows'][0]['status'], 'Normal')
        self.assertIsNone(data['rows'][0]['decision'])
        path = Path(self.folder.name) / 'green-recommendation.pdf'
        export_pdf(data, path)
        document = pdfium.PdfDocument(str(path))
        texts = []
        try:
            for index in range(len(document)):
                page = document[index]
                textpage = page.get_textpage()
                try:
                    texts.append(textpage.get_text_range())
                finally:
                    textpage.close()
                    page.close()
        finally:
            document.close()
        shown = ' '.join(' '.join(texts).split())
        expected = ' '.join(data['rows'][0]['recommendation']['next_action'].split())
        self.assertIn(expected, shown)

    def test_unmapped_material_does_not_claim_a_demo_destination(self):
        key = Sessions(self.db).start("Petugas Uji", "Area A")
        Sessions(self.db).add_code(key, "DEMO-0001")
        advice = report_data(self.db, key)["rows"][0]["recommendation"]
        self.assertNotIn("source", advice)
        self.assertNotIn("Site", advice["next_action"])
        self.assertNotIn("Customer", advice["next_action"])


if __name__ == "__main__":
    unittest.main()
