"""Report reads and exports always use session snapshots, never current inventory."""
from datetime import date,datetime
from pathlib import Path
from contextlib import contextmanager
from xml.sax.saxutils import escape
import os
import tempfile
from .sessions import Sessions
from .recommendations import RecommendationService, recommend_for

STATUSES=('Normal','Warning','Melewati batas')


def decision_target(decision):
    """Human-readable destination recorded by the officer, if one was chosen."""
    if not decision:
        return ''
    kind, code = decision.get('target_kind'), decision.get('target_code')
    if kind == 'site' and code:
        return f'Site {code}'
    if kind == 'customer' and code:
        return f'Customer {code}'
    return ''


def history(db,query='',status='',date_from='',date_to=''):
    for value in (date_from,date_to):
        if value:
            try:
                if date.fromisoformat(value).isoformat()!=value:raise ValueError()
            except ValueError as exc:raise ValueError('Tanggal filter harus YYYY-MM-DD.') from exc
    if date_from and date_to and date_from>date_to:raise ValueError('Tanggal awal tidak boleh setelah tanggal akhir.')
    if status and status not in ('Draft','Selesai'):raise ValueError('Status laporan tidak valid.')
    result=[]
    for session in Sessions(db).list():
        day=session['started_at'][:10]
        text=f"CHK-{session['id']:06d} {session['officer']} {session['area']}".casefold()
        if (query.strip().casefold() in text and (not status or session['status']==status)
                and (not date_from or day>=date_from) and (not date_to or day<=date_to)):
            result.append(session)
    return result


def report_data(db,key):
    service=Sessions(db)
    session=service.get(key)
    rows=service.results(key)
    decisions=RecommendationService(db)
    immutable_fields=('id','physical_condition','choice','reason','officer','decided_at',
                      'age_at_decision','aging_status','recommendation_code','recommendation_title',
                      'recommendation_reason','recommendation_next_action','recommendation_priority',
                      'target_kind','target_code')
    for row in rows:
        checked_day=date.fromisoformat(row['checked_at'][:10])
        row['recommendation']=row.get('recommendation_at_check') or recommend_for(
            row,normal_max=session['normal_max'],age_limit=session['age_limit'],today=checked_day)
        recorded=decisions.for_result(key,row['id'])
        row['decision']={field:recorded[field] for field in immutable_fields} if recorded else None
    return dict(session=session,rows=rows,counts={s:sum(r['status']==s for r in rows) for s in STATUSES},
                generated_at=datetime.now().astimezone().isoformat(timespec='seconds'))


@contextmanager
def atomic_file(path):
    path=Path(path)
    with tempfile.NamedTemporaryFile(dir=path.parent,suffix=path.suffix,delete=False) as file:temporary=Path(file.name)
    try:
        yield temporary
        os.replace(temporary,path)
    finally:
        if temporary.exists():temporary.unlink()


def export_excel(data,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    session,rows=data['session'],data['rows']
    book=Workbook()
    summary=book.active
    summary.title='Ringkasan'
    pairs=[('LAPORAN PENGECEKAN WAREHOUSE','DATA DEMO'),('Nomor',f"CHK-{session['id']:06d}"),
           ('Status',session['status']),('Petugas',session['officer']),('Area',session['area']),
           ('Mulai',session['started_at']),('Selesai',session['finished_at'] or 'Belum selesai'),
           ('Normal maksimum (hari)',session['normal_max']),('Batas umur (hari)',session['age_limit']),
           ('Total batch',len(rows)),*data['counts'].items(),('Catatan sesi',session['note']),
           ('Diekspor',data['generated_at']),('Cakupan','Hanya barang dalam sesi. Status umur bukan hasil pemeriksaan fisik.')]
    for pair in pairs:summary.append(pair)
    sheet=book.create_sheet('Hasil pengecekan')
    headers=('No','Waktu cek','Batch ID','Material','Deskripsi','Plant','Storage Location','Lokasi','GR Date','Stok','Umur (hari)','Status','Catatan',
             'Rekomendasi saat cek','Kondisi fisik','Keputusan petugas','Alasan keputusan','Diputuskan oleh','Waktu keputusan','Tujuan tindakan')
    sheet.append(headers)
    for number,row in enumerate(rows,1):
        decision=row.get('decision') or {}
        sheet.append([number,*[row[k] for k in ('checked_at','id','material','description','plant','storage_location','location_code','received_date','stock','age','status','note')],
                      row['recommendation']['next_action'],decision.get('physical_condition',''),
                      decision.get('choice',''),decision.get('reason',''),decision.get('officer',''),
                      decision.get('decided_at',''),decision_target(decision)])
    palette={'Normal':'DCF4E6','Warning':'FFF0C4','Melewati batas':'FFDDE0'}
    for ws in book:
        for row in ws:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='11283F')
            cell.font=Font(bold=True,color='FFFFFF')
        ws.row_dimensions[1].height=30
        ws.freeze_panes='A2'
    summary.column_dimensions['A'].width=32
    summary.column_dimensions['B'].width=100
    summary.row_dimensions[14].height=min(409,20*max(2,len(session['note'])//90+session['note'].count('\n')+1))
    widths=(7,28,42,23,36,14,20,22,16,10,12,22,65,75,20,27,70,22,28,24)
    for column,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(column)].width=width
    for index,row in enumerate(rows,2):
        sheet.cell(index,12).fill=PatternFill('solid',fgColor=palette[row['status']])
        decision=row.get('decision') or {}
        longest=max(len(row['note'])//60+row['note'].count('\n')+1,
                    len(row['recommendation']['next_action'])//70+1,
                    len(decision.get('reason',''))//70+1)
        sheet.row_dimensions[index].height=min(409,18*max(3,longest))
    sheet.auto_filter.ref=sheet.dimensions
    sheet.print_title_rows='1:1'
    sheet.page_setup.orientation='landscape'
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.fitToWidth=1
    sheet.page_setup.fitToHeight=0
    try:
        with atomic_file(path) as temporary:book.save(temporary)
    finally:book.close()


def export_pdf(data,path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    regular,bold='Helvetica','Helvetica-Bold'
    font_root=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'
    if (font_root/'arial.ttf').exists():
        if 'WarehouseArial' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('WarehouseArial',str(font_root/'arial.ttf')))
            pdfmetrics.registerFont(TTFont('WarehouseArialBold',str(font_root/'arialbd.ttf')))
        regular,bold='WarehouseArial','WarehouseArialBold'
    normal=ParagraphStyle('body',fontName=regular,fontSize=9,leading=13,spaceAfter=6,wordWrap='LTR')
    small=ParagraphStyle('cell',parent=normal,fontSize=8,leading=11,spaceAfter=0)
    header=ParagraphStyle('head',parent=small,fontName=bold,textColor=colors.white)
    title=ParagraphStyle('title',parent=normal,fontName=bold,fontSize=20,leading=25,textColor=colors.HexColor('#11283F'))
    def p(value,style=normal):return Paragraph(escape(str(value)).replace('\n','<br/>'),style)
    session,rows=data['session'],data['rows']
    report_id=f"CHK-{session['id']:06d}"
    story=[p('LAPORAN PENGECEKAN WAREHOUSE',title),p(f"{report_id} | {session['status'].upper()} | DATA DEMO"),
           p(f"Petugas: {session['officer']}  |  Area: {session['area']}"),
           p(f"Mulai: {session['started_at']}  |  Selesai: {session['finished_at'] or 'Belum selesai'}"),
           p(f"Normal: 0–{session['normal_max']} hari  |  Warning: {session['normal_max']+1}–{session['age_limit']} hari  |  Melewati batas: >{session['age_limit']} hari"),
           p(f"Total: {len(rows)} batch  |  "+'  |  '.join(f'{s}: {data["counts"][s]}' for s in STATUSES)),
           p('Cakupan: hanya barang dalam sesi ini. Status umur bukan hasil pemeriksaan fisik.')]
    if session['status']=='Draft':story.append(p('DRAFT — sesi belum selesai; hasil dapat bertambah sampai sesi dikunci.'))
    if session['note']:story.extend([p('Catatan sesi:',ParagraphStyle('bold',parent=normal,fontName=bold)),p(session['note'])])
    story.append(Spacer(1,10))
    headers=('No','Batch / barang','Lokasi / plant','GR Date / waktu cek','Stok','Umur','Status','Saran / keputusan')
    table_rows=[[p(h,header) for h in headers]]
    for i,row in enumerate(rows,1):
        values=(i,f"{row['id']}\n{row['material']}\n{row['description']}",
                f"{row['location_code']}\n{row['plant']} / {row['storage_location']}",
                f"{row['received_date']}\nCek: {row['checked_at']}",row['stock'],f"{row['age']} hari",row['status'],
                (f"{row['recommendation']['title']}\n"
                 f"{row['decision']['choice']}"
                 f"{' → ' + decision_target(row['decision']) if decision_target(row['decision']) else ''}"
                 if row.get('decision') else f"{row['recommendation']['title']}\nBelum ada keputusan"))
        table_rows.append([p(v,small) for v in values])
    if not rows:story.append(p('Belum ada hasil pengecekan pada sesi ini.'))
    table=LongTable(table_rows,colWidths=[30,205,135,130,40,48,90,100],repeatRows=1,hAlign='LEFT')
    rules=[('BACKGROUND',(0,0),(-1,0),colors.HexColor('#11283F')),('VALIGN',(0,0),(-1,-1),'TOP'),
           ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
           ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),
           ('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#D4DDE5'))]
    palette={'Normal':'#DCF4E6','Warning':'#FFF0C4','Melewati batas':'#FFDDE0'}
    for i,row in enumerate(rows,1):rules.append(('BACKGROUND',(6,i),(6,i),colors.HexColor(palette[row['status']])))
    table.setStyle(TableStyle(rules))
    story.append(table)
    detail_title=ParagraphStyle('note',parent=normal,fontName=bold)
    for i,row in enumerate(rows,1):
        decision=row.get('decision')
        if (row['status']=='Normal' and not row['note'] and not decision
                and not row['recommendation'].get('options')):continue
        advice_label=('Rekomendasi saat cek (data simulasi): '
                      if row['recommendation'].get('source')=='Simulasi'
                      else 'Saran berdasarkan umur saat cek: ')
        story.extend([Spacer(1,12),p(f"Tindak lanjut #{i} — {row['id']} / {row['material']}",detail_title),
                      p(advice_label+row['recommendation']['next_action'])])
        if row['note']:story.append(p('Catatan hasil: '+row['note']))
        if decision:
            story.extend([p(f"Kondisi fisik: {decision['physical_condition']} | Keputusan: {decision['choice']}"),
                          p(f"Alasan petugas: {decision['reason']}"),
                          p(f"Dicatat oleh {decision['officer']} pada {decision['decided_at']}")])
            if decision_target(decision):
                story.append(p('Tujuan tindakan: '+decision_target(decision)))
    story.extend([Spacer(1,12),p('Diekspor: '+data['generated_at'],small)])
    def footer(canvas,doc):
        canvas.saveState()
        canvas.setFont(regular,8)
        canvas.setFillColor(colors.HexColor('#526476'))
        canvas.drawString(28,17,f'{report_id} | {session["status"]} | Warehouse Check')
        canvas.drawRightString(landscape(A4)[0]-28,17,f'Halaman {doc.page}')
        canvas.restoreState()
    with atomic_file(path) as temporary:
        doc=SimpleDocTemplate(str(temporary),pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=25,bottomMargin=35,
                              title=f'Laporan {report_id}',author='Warehouse Check')
        doc.build(story,onFirstPage=footer,onLaterPages=footer)
