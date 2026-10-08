"""Inventory XLSX interchange with preview and atomic, revalidated commits."""
from datetime import date, datetime
from pathlib import Path
import os
import re
import tempfile
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .domain import legacy_code_to_new
from .inventory import Inventory
from .locations import Locations

HEADERS = ('Batch ID', 'Material', 'Material Description', 'Plant', 'Storage Location', 'Location Code', 'GR Date', 'Total Stock')
FIELDS = ('id', 'material', 'description', 'plant', 'storage_location', 'location_code', 'received_date', 'stock')
MAX_ROWS = 5000


def read_inventory(path):
    if Path(path).suffix.lower() != '.xlsx':
        raise ValueError('Gunakan file .xlsx sesuai template.')
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Inventory' not in book.sheetnames:
            raise ValueError('Sheet Inventory tidak ditemukan. Gunakan template aplikasi.')
        sheet = book['Inventory']
        # Some valid XLSX files omit worksheet dimensions. In read-only mode
        # openpyxl then reports max_row=None; count rows while iterating too.
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
            raise ValueError(f'Maksimal {MAX_ROWS} baris per impor. Pisahkan file menjadi beberapa bagian.')
        iterator = sheet.iter_rows()
        headers = [str(c.value or '').strip() for c in next(iterator, [])]
        if len(headers) != len(set(headers)) or any(h not in headers for h in HEADERS):
            raise ValueError('Header hilang atau duplikat. Kolom wajib: ' + ', '.join(HEADERS))
        indexes = [headers.index(h) for h in HEADERS]
        result = []
        for number, cells in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                raise ValueError(f'Maksimal {MAX_ROWS} baris per impor. Pisahkan file menjadi beberapa bagian.')
            values, issues = {}, []
            for header, key, index in zip(HEADERS, FIELDS, indexes):
                cell = cells[index] if index < len(cells) else None
                value = cell.value if cell is not None else None
                if cell is not None and cell.data_type in ('f', 'e'):
                    issues.append(f'{header}: rumus/error Excel tidak didukung; gunakan nilai biasa.')
                if key == 'received_date' and isinstance(value, (date, datetime)):
                    value = value.strftime('%Y-%m-%d')
                if key == 'stock' and isinstance(value, float) and value.is_integer():
                    value = int(value)
                if key in ('id','material','plant','storage_location','location_code') and value is not None and not isinstance(value,str):
                    issues.append(f'{key}: kode harus berupa teks agar nol di depan tidak hilang.')
                values[key] = '' if value is None else str(value).strip()
            if not any(values.values()):
                continue
            result.append(dict(row=number, values=values, issues=issues))
        if not result:
            raise ValueError('Sheet Inventory belum berisi data.')
        return result
    finally:
        book.close()


def fingerprint(values):
    return tuple(str(values[k]).strip().casefold() for k in FIELDS if k != 'id')


def preview(db, source):
    inventory = Inventory(db)
    existing = [dict(r) for r in db.execute('SELECT * FROM batches')]
    existing_ids = {r['id'] for r in existing}
    existing_data = {fingerprint(r) for r in existing}
    seen_ids, seen_data, result = set(), set(), []
    for record in source:
        values = dict(record['values'])
        errors = list(record.get('issues', []))
        notices = []
        key = values.get('id', '')
        if (re.fullmatch(r'\d{4}-\d{2}-\d{2}', values.get('location_code', ''))
                and not values.get('stock', '')):
            errors.append('Kolom tampak bergeser. Jika Batch ID belum ada, biarkan kolom A kosong; isi Material di B, Location Code di F, GR Date di G, dan Total Stock di H.')
        elif not values.get('stock', ''):
            errors.append('Total Stock (kolom H) kosong. Periksa urutan kolom pada baris ini.')
        if key and not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,99}', key):
            errors.append('Batch ID hanya huruf, angka, - atau _, maksimal 100 karakter.')
        if key and (key in existing_ids or key in seen_ids):
            errors.append('Batch ID duplikat di file/database; gunakan Edit barang untuk pembaruan.')
        old_code = str(values.get('location_code', '')).strip().upper()
        if re.fullmatch(r'[A-Z]-[A-Z][0-9]{2}-\d{2}-\d{2}', old_code):
            try:
                translated = legacy_code_to_new(old_code)
            except ValueError as exc:
                errors.append(str(exc))
            else:
                values['location_code'] = translated
                notices.append(f"Kode lokasi lama ditafsirkan sebagai {translated} (baris 01).")
        try:
            clean = inventory.validate(values)
            signature = fingerprint(dict(clean,id=key))
            if not key and (signature in existing_data or signature in seen_data):
                errors.append('Data identik terdeteksi. Jika penerimaan berbeda, isi Batch ID unik secara eksplisit.')
            seen_data.add(signature)
        except ValueError as exc:
            errors.append(str(exc))
            clean = values
        if key:seen_ids.add(key)
        result.append(dict(row=record['row'], values=dict(clean,id=key), errors=errors, notices=notices))
    return result


def commit_import(db, source):
    if not source:
        raise ValueError('Tidak ada data untuk disimpan.')
    # Lock before revalidation, so another connection cannot insert between check and commit.
    db.execute('BEGIN IMMEDIATE')
    try:
        checked = preview(db, source)
        if any(r['errors'] for r in checked):
            raise ValueError('Impor dibatalkan: ada kesalahan/duplikasi. Muat ulang preview untuk melihat detail.')
        keys = []
        for row in checked:
            values = dict(row['values'])
            values['id'] = values['id'] or 'BAT-' + uuid4().hex.upper()
            db.execute('''INSERT INTO batches
                (id,material,description,plant,storage_location,location_code,received_date,stock)
                VALUES (:id,:material,:description,:plant,:storage_location,:location_code,:received_date,:stock)''', values)
            keys.append(values['id'])
        db.commit()
        return keys
    except Exception:
        db.rollback()
        raise


def write_inventory(path, db, template=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Inventory'
    headers = HEADERS if template else HEADERS + ('Aging (days)', 'Status', 'As of')
    sheet.append(headers)
    if not template:
        for row in Inventory(db).list_batches():
            sheet.append([row[k] for k in FIELDS] + [row['age'],row['status'],date.today().isoformat()])
    else:
        # Empty input row with text formatting, never a sample that could be imported accidentally.
        sheet.append([''] * len(HEADERS))
    locations = book.create_sheet('Lokasi')
    locations.append(['Location Code', 'Area', 'Rak', 'Kolom', 'Baris', 'Tingkat'])
    for row in Locations(db).list():
        locations.append([row[k] for k in ('code','area','rack','column','row','level')])
    guide = book.create_sheet('Petunjuk')
    for line in ['Warehouse Check — template inventory / data demo',
                 'Isi sheet Inventory; jangan mengganti nama header.',
                 'Batch ID boleh kosong: aplikasi membuat ID. ID terisi harus unik; impor tidak mengubah batch lama.',
                 'Kolom kode harus Teks. GR Date berupa tanggal Excel atau YYYY-MM-DD, bukan tanggal masa depan.',
                 'Urutan kolom: A Batch ID, B Material, C Material Description, D Plant, E Storage Location, F Location Code, G GR Date, H Total Stock.',
                 'Jika Batch ID belum ada, biarkan sel kolom A kosong. Jangan menggeser data mulai dari kolom B ke A.',
                 'Total Stock bilangan bulat 0–999999999. Semua kolom selain Batch ID wajib.',
                 'Location Code harus terdaftar. Lihat sheet Lokasi atau tambah melalui Kelola Lokasi.',
                 'Kode lokasi baru: A-B01-K03-R02-T02 = Area A, Rak B01, Kolom 03, Baris 02, Tingkat 02.',
                 'Kode empat bagian dari template lama diterjemahkan menjadi baris 01 saat preview; periksa hasilnya sebelum menyimpan.',
                 'Maksimal 5000 baris. Rumus Excel ditolak. Preview harus bebas kesalahan sebelum simpan.',
                 'Aging, Status, As of pada ekspor adalah nilai saat ekspor dan diabaikan saat impor.',
                 'Ekspor memuat seluruh inventory, tidak mengikuti filter layar.']:
        guide.append([line])
    for ws in book:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='11283F')
            cell.font = Font(color='FFFFFF',bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment = Alignment(vertical='top',wrap_text=True)
        for col in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width = 26
    sheet.column_dimensions['A'].width = 42
    sheet.column_dimensions['C'].width = 36
    for row in sheet.iter_rows(min_row=2):
        for cell in row[:7]:cell.number_format='@'
    guide.column_dimensions['A'].width = 110
    temporary=None
    try:
        with tempfile.NamedTemporaryFile(dir=Path(path).parent,suffix='.xlsx',delete=False) as file:
            temporary=Path(file.name)
        book.save(temporary)
        os.replace(temporary,path)
    finally:
        book.close()
        if temporary is not None and temporary.exists():temporary.unlink()
