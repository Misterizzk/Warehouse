"""Build a new portable Windows folder without overwriting user data."""
import argparse
from contextlib import closing
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]


def create_clean_install_database(seed, target):
    """Copy demo master data to a fresh schema, without previous check sessions."""
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from warehouse.database import connect

    seed = Path(seed).resolve()
    target = Path(target)
    with closing(sqlite3.connect(seed.as_uri() + '?mode=ro', uri=True)) as source:
        locations = source.execute('''SELECT code,area,rack,"column","row",level
            FROM locations ORDER BY code''').fetchall()
        batches = source.execute('''SELECT id,material,description,plant,storage_location,
            location_code,received_date,stock FROM batches ORDER BY id''').fetchall()
        site_demand = source.execute('''SELECT material,site,demand_prev,demand_now,
            open_demand,site_stock FROM site_demand ORDER BY material,site''').fetchall()
        compatibility = source.execute('''SELECT material,unit_model
            FROM material_unit_compatibility ORDER BY material,unit_model''').fetchall()
        customer_demand = source.execute('''SELECT customer,unit_model,open_demand
            FROM customer_demand ORDER BY customer,unit_model''').fetchall()
        seeded = source.execute("SELECT value FROM metadata WHERE key='seed_date'").fetchone()
    installed = connect(target)
    try:
        with installed:
            installed.executemany('''INSERT INTO locations
                (code,area,rack,"column","row",level) VALUES(?,?,?,?,?,?)''', locations)
            installed.executemany('''INSERT INTO batches
                (id,material,description,plant,storage_location,location_code,received_date,stock)
                VALUES(?,?,?,?,?,?,?,?)''', batches)
            installed.executemany('''INSERT INTO site_demand
                (material,site,demand_prev,demand_now,open_demand,site_stock)
                VALUES(?,?,?,?,?,?)''', site_demand)
            installed.executemany('''INSERT INTO material_unit_compatibility
                (material,unit_model) VALUES(?,?)''', compatibility)
            installed.executemany('''INSERT INTO customer_demand
                (customer,unit_model,open_demand) VALUES(?,?,?)''', customer_demand)
            installed.execute("INSERT INTO metadata(key,value) VALUES('seed_date',?)",
                              (seeded[0] if seeded else datetime.now().date().isoformat(),))
        if installed.execute('PRAGMA foreign_key_check').fetchall():
            raise ValueError('Lokasi barang pada data uji tidak valid.')
    finally:
        installed.close()


def create_clean_install_workbook(database_path, target):
    """Export demo inventory and transparent scenarios, omitting staff history."""
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    from warehouse.database import connect
    from warehouse.inventory_excel import write_inventory
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    installed = connect(database_path)
    try:
        write_inventory(target, installed)
        book = load_workbook(target)
        try:
            guide = book['Petunjuk']
            guide.append(['Sheet Demand_Site_Simulasi, Kecocokan_Unit, dan Customer_Simulasi '
                          'menjelaskan data rekaan rekomendasi A/B/C. Hanya Inventory dibaca '
                          'saat impor; mengubah sheet simulasi tidak mengubah rekomendasi aplikasi.'])
            guide.auto_filter.ref = guide.dimensions
            for name, table in (
                ('Demand_Site_Simulasi', 'site_demand'),
                ('Kecocokan_Unit', 'material_unit_compatibility'),
                ('Customer_Simulasi', 'customer_demand'),
            ):
                columns = [row['name'] for row in installed.execute(f'PRAGMA table_info("{table}")')]
                sheet = book.create_sheet(name)
                sheet.append(columns)
                for row in installed.execute(f'SELECT * FROM "{table}" ORDER BY 1,2'):
                    sheet.append(tuple(row))
                sheet.freeze_panes = 'A2'
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.fill = PatternFill('solid', fgColor='11283F')
                    cell.font = Font(color='FFFFFF', bold=True)
                for row in sheet.iter_rows(min_row=2):
                    for cell in row:
                        if isinstance(cell.value, str):
                            cell.data_type = 's'
                for index in range(1, sheet.max_column + 1):
                    sheet.column_dimensions[get_column_letter(index)].width = 28
            book.save(target)
        finally:
            book.close()
    finally:
        installed.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', default='Aplikasi_Warehouse_Check')
    args = parser.parse_args()
    if sys.platform != 'win32':
        raise SystemExit('Build ini harus dijalankan pada Windows 64-bit.')
    output = (ROOT / args.output).resolve()
    if not output.is_relative_to(ROOT) or output == ROOT:
        raise SystemExit('Folder hasil harus berada di dalam folder proyek.')
    if output.exists():
        raise SystemExit('Folder hasil sudah ada. Gunakan --output nama_baru agar database pengguna tidak tertimpa.')
    seed = ROOT / 'data_uji' / 'warehouse_demo_seed.db'
    for source in (seed, ROOT / 'warehouse' / 'package_check.py'):
        if not source.is_file():
            raise SystemExit(f'File belum tersedia: {source}')
    # Fail before packaging if Tcl cannot initialize. PyInstaller otherwise
    # silently excludes Tkinter and creates an executable that cannot start.
    subprocess.run([sys.executable, '-c',
                    'import tkinter as tk; root=tk.Tk(); root.withdraw(); root.destroy()'],
                   cwd=ROOT, check=True)
    stage = ROOT / '.build' / ('package_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
    stage.mkdir(parents=True, exist_ok=False)
    if not stage.resolve().is_relative_to(ROOT / '.build'):
        raise SystemExit('Folder build tidak valid.')
    icon = ROOT / 'packaging' / 'warehouse.ico'
    subprocess.run([
        sys.executable, '-m', 'PyInstaller', '--onedir', '--windowed', '--noupx',
        '--name', 'WarehouseCheck', '--contents-directory', '_internal',
        '--icon', str(icon), '--version-file', str(ROOT / 'packaging' / 'version_info.txt'),
        '--distpath', str(stage / 'dist'), '--workpath', str(stage / 'work'),
        '--specpath', str(stage), '--add-data', str(ROOT / 'assets') + os.pathsep + 'assets',
        '--collect-all', 'pypdfium2_raw', '--collect-data', 'reportlab',
        '--exclude-module', 'pandas', '--exclude-module', 'matplotlib',
        '--exclude-module', 'pytest', '--exclude-module', 'IPython',
        str(ROOT / 'warehouse_app.py'),
    ], cwd=ROOT, check=True)
    candidate = stage / 'dist' / 'WarehouseCheck'
    (candidate / 'data').mkdir()
    create_clean_install_database(seed, candidate / 'data' / 'warehouse_demo.db')
    create_clean_install_workbook(candidate / 'data' / 'warehouse_demo.db',
                                  candidate / 'Data_Uji_Warehouse.xlsx')
    shutil.copy2(ROOT / 'packaging' / 'MULAI_DI_SINI.txt', candidate / 'MULAI_DI_SINI.txt')
    frozen_report = stage / 'package_verification.json'
    # Different working directory verifies portable path handling.
    result = subprocess.run([str(candidate / 'WarehouseCheck.exe'), '--verify-package', str(frozen_report)],
                            cwd=stage, timeout=150)
    if result.returncode or not frozen_report.exists():
        raise SystemExit(f'Uji executable gagal. Lihat folder: {stage}')
    report = json.loads(frozen_report.read_text(encoding='utf-8-sig'))
    if not report.get('ok'):
        raise SystemExit(f'Uji executable gagal: {report}')
    manifest = {
        'built_at': datetime.now().astimezone().isoformat(timespec='seconds'),
        'python_version': sys.version,
        'type': 'Windows x64 portable folder',
        'source_data_manifest': json.loads((ROOT / 'data_uji' / 'manifest.json').read_text(encoding='utf-8-sig')),
        'installed_database': {'inventory_batches': 200, 'check_sessions': 0,
                               'check_results': 0, 'recommendation_decisions': 0},
        'exe_sha256': hashlib.sha256((candidate / 'WarehouseCheck.exe').read_bytes()).hexdigest(),
        'initial_database_sha256': hashlib.sha256((candidate / 'data' / 'warehouse_demo.db').read_bytes()).hexdigest(),
        'verification': report,
    }
    (candidate / 'manifest_paket.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    shutil.copytree(candidate, output)
    print(f'Selesai: {output / "WarehouseCheck.exe"}', flush=True)
    print(f'Ukuran: {sum(p.stat().st_size for p in output.rglob("*") if p.is_file()) / 1024**2:.1f} MB', flush=True)


if __name__ == '__main__':
    main()
