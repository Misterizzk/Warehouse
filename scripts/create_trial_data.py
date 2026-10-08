"""Build an importable trial workbook and seed database without editing sources.

Run from any working directory:
    python scripts/create_trial_data.py --today 2026-09-27
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
import sqlite3
import sys
import tempfile

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from warehouse.database import connect
from warehouse.domain import aging, location_code
from warehouse.inventory_excel import FIELDS, HEADERS, commit_import, preview, read_inventory

LOCATION_FIELDS = ('code', 'area', 'rack', 'column', 'row', 'level')
ADDED_RACKS = [('C', 'C02'), ('C', 'C03'), ('C', 'C04'), ('D', 'D01'), ('D', 'D02')]
EXTRA_RACKS = [('A', 'B01'), ('A', 'B02'), ('A', 'B03'), ('B', 'B01'),
               ('B', 'B02'), ('B', 'B03'), ('C', 'C01')]
PARTS = [
    ('HYD-001', 'HYDRAULIC PUMP ASSY'), ('HYD-002', 'HOSE HIGH PRESSURE 1/2 INCH'),
    ('FLT-001', 'OIL FILTER ELEMENT'), ('FLT-002', 'AIR FILTER PRIMARY'),
    ('SEA-001', 'SEAL KIT CYLINDER'), ('BRG-001', 'ROLLER BEARING'),
    ('ELC-001', 'TEMPERATURE SENSOR'), ('ELC-002', 'RELAY 24V'),
    ('BRK-001', 'BRAKE PAD SET'), ('ENG-001', 'FUEL INJECTOR'),
    ('FAS-001', 'BOLT M12 X 40'), ('FAS-002', 'NUT M12'),
    ('TRN-001', 'GEAR TRANSMISSION'), ('CAB-001', 'WORK LAMP LED'),
    ('0000123405', 'O-RING 24 MM'), ('0000000786', 'WASHER FLAT 12 MM'),
]


def readonly_database(path: Path):
    db = sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)
    db.row_factory = sqlite3.Row
    return db


def table_rows(db, table):
    # Table identifiers below come only from this module/schema, never user text.
    return [dict(row) for row in db.execute(f'SELECT * FROM "{table}" ORDER BY rowid')]


def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_legacy_inventory(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book['Inventory']
        sheet.reset_dimensions()  # The original workbook has no reliable dimensions.
        iterator = iter(sheet.values)
        headers = list(next(iterator))
        records = []
        for number, cells in enumerate(iterator, 2):
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            record = dict(zip(headers, list(cells) + [None] * (len(headers) - len(cells))))
            records.append((number, record))
        return headers, records
    finally:
        book.close()


def _date_text(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    return str(value).strip()


def _insert_batch(db, values):
    db.execute('''INSERT INTO batches
        (id, material, description, plant, storage_location, location_code, received_date, stock)
        VALUES (:id,:material,:description,:plant,:storage_location,:location_code,:received_date,:stock)''', values)


def _add_synthetic(db, key, material, description, code, age, stock, today,
                   provenance, scenario, plant='BHT', storage='PBHT'):
    values = dict(id=key, material=material, description=description, plant=plant,
                  storage_location=storage, location_code=code,
                  received_date=(today-timedelta(days=age)).isoformat(), stock=stock)
    _insert_batch(db, values)
    provenance.append([key, 'Simulasi tambahan', 'Generator data uji', '', '', code, scenario])


def _seed_recommendation_scenarios(db, today):
    """Record explicit, deterministic A/B/C demand for the demo materials."""
    site_rows = [
        # Normal, rising demand at A; Normal, no shortage; Warning, customer.
        ('MAT-REC-001','A',3,9,10,1), ('MAT-REC-001','B',3,3,2,8),
        ('MAT-REC-001','C',2,2,0,4),
        ('MAT-REC-002','A',4,4,0,12), ('MAT-REC-002','B',1,2,1,5),
        ('MAT-REC-002','C',2,2,0,6),
        ('MAT-REC-003','A',3,3,1,5), ('MAT-REC-003','B',2,2,0,4),
        ('MAT-REC-003','C',1,1,0,3),
        # Critical stock: customer A, Site B, and both customer A/Site C.
        ('MAT-REC-004','A',2,2,0,5), ('MAT-REC-004','B',2,2,1,4),
        ('MAT-REC-004','C',1,1,0,3),
        ('MAT-REC-005','A',2,2,0,6), ('MAT-REC-005','B',3,12,15,1),
        ('MAT-REC-005','C',2,2,1,7),
        ('MAT-REC-006','A',2,2,0,5), ('MAT-REC-006','B',2,2,1,6),
        ('MAT-REC-006','C',4,14,12,1),
    ]
    compatibility = [
        ('MAT-REC-003','UNIT-CUSTOMER-A-WARNING'),
        ('MAT-REC-004','UNIT-CUSTOMER-A-CRITICAL'),
        ('MAT-REC-006','UNIT-CUSTOMER-A-CHOICE'),
        ('FLT-001','UNIT-CUSTOMER-B'),
        ('HYD-001','UNIT-CUSTOMER-C'),
    ]
    customer_rows = [
        ('A','UNIT-CUSTOMER-A-WARNING',4),
        ('A','UNIT-CUSTOMER-A-CRITICAL',3),
        ('A','UNIT-CUSTOMER-A-CHOICE',5),
        ('B','UNIT-CUSTOMER-B',2),
        ('C','UNIT-CUSTOMER-C',4),
    ]
    # The other synthetic parts also have a mix of shortages and stable sites,
    # so the dashboard can demonstrate recommendations beyond six showcase rows.
    for index, (material, _description) in enumerate(PARTS):
        site_rows.extend([
            (material,'A',3+index%3,5+index%5,5+index%4,1+index%3),
            (material,'B',2+index%4,2+index%4,1+index%3,5+index%3),
            (material,'C',1+index%3,4+index%5,4+index%3,1+index%2),
        ])
    materials = [row[0] for row in db.execute('SELECT DISTINCT material FROM batches ORDER BY material')]
    for material in materials:
        # Generate only during trial-data creation. A material imported later
        # into the app has no fabricated destination until data is supplied.
        key = int.from_bytes(hashlib.sha256(material.encode('utf-8')).digest()[:4], 'big')
        hot_site = key % 3
        for index, site in enumerate(('A','B','C')):
            previous = 2 + (key >> (index*3)) % 5
            if index == hot_site:
                site_rows.append((material,site,previous,previous+3+(key%4),
                                  5+(key%5),1+(key%2)))
            else:
                site_rows.append((material,site,previous,previous,
                                  1+(key%2),5+(key%4)))
    db.executemany('''INSERT OR IGNORE INTO site_demand
        (material,site,demand_prev,demand_now,open_demand,site_stock)
        VALUES(?,?,?,?,?,?)''', site_rows)
    db.executemany('''INSERT OR IGNORE INTO material_unit_compatibility
        (material,unit_model) VALUES(?,?)''', compatibility)
    db.executemany('''INSERT OR IGNORE INTO customer_demand
        (customer,unit_model,open_demand) VALUES(?,?,?)''', customer_rows)
    for material in materials:
        if material == 'MAT-REC-005':  # Curated red/site-only PIC example.
            continue
        has_aging_batch = any(
            aging(row['received_date'], today=today)[0] > 30
            for row in db.execute('''SELECT received_date FROM batches
                WHERE material=? AND stock>0''', (material,)))
        if not has_aging_batch:
            continue
        has_customer_need = db.execute('''SELECT 1
            FROM material_unit_compatibility c JOIN customer_demand d
              ON d.unit_model=c.unit_model
            WHERE c.material=? AND d.open_demand>0 LIMIT 1''', (material,)).fetchone()
        if has_customer_need:
            continue
        key = int.from_bytes(hashlib.sha256(material.encode('utf-8')).digest()[:4], 'big')
        customer = ('A','B','C')[key % 3]
        unit_model = 'UNIT-SIM-' + material
        db.execute('''INSERT OR IGNORE INTO material_unit_compatibility
            (material,unit_model) VALUES(?,?)''', (material,unit_model))
        db.execute('''INSERT OR IGNORE INTO customer_demand
            (customer,unit_model,open_demand) VALUES(?,?,?)''',
            (customer,unit_model,1+key%6))


def _sheet(book, name, headers, rows):
    ws = book.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='14324B')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 31
        for cell in row:
            cell.font = Font(name='Arial', size=10, color='20384C')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EFF5FA')
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Literal source text, including an initial '='.
                cell.number_format = '@'
            elif isinstance(cell.value, (date, datetime)):
                cell.number_format = 'yyyy-mm-dd'
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
    for col in range(1, ws.max_column+1):
        ws.column_dimensions[get_column_letter(col)].width = 23
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    return ws


def _dict_sheet(book, name, rows, fallback_headers):
    headers = list(rows[0]) if rows else list(fallback_headers)
    return _sheet(book, name, headers, ([r.get(k) for k in headers] for r in rows))


def _workbook(path, db, originals, legacy_headers, legacy_records, old_history,
              mapping, provenance, today):
    book = Workbook()
    book.remove(book.active)
    inventory = table_rows(db, 'batches')
    locations = table_rows(db, 'locations')
    ws = _sheet(book, 'Inventory', HEADERS, ([r[k] for k in FIELDS] for r in inventory))
    for col, width in {'A': 43, 'B': 23, 'C': 51, 'D': 13, 'E': 22, 'F': 29, 'G': 17, 'H': 16}.items():
        ws.column_dimensions[col].width = width
    for cell in ws[1]:
        cell.comment = Comment('Header ini digunakan oleh impor aplikasi. Jangan diganti atau digeser.', 'Warehouse Check')
    ws['A1'].comment = Comment('ID harus unik. File ini adalah salinan data pada aplikasi uji; impor ke aplikasi yang sudah berisi data yang sama akan ditolak sebagai duplikat.', 'Warehouse Check')
    _sheet(book, 'Lokasi', ['Location Code', 'Area', 'Rak', 'Kolom', 'Baris', 'Tingkat'],
           ([r[k] for k in LOCATION_FIELDS] for r in locations))
    guide = [
        ['Tujuan', 'Satu workbook berisi inventory untuk uji coba, sumber asli, pemetaan lokasi lama, dan seluruh riwayat yang tersedia.'],
        ['Tanggal acuan skenario', today.isoformat() + '. Data baru dibuat relatif terhadap tanggal ini; tanggal tidak digeser otomatis setelah file dibuat.'],
        ['Sumber lama tetap utuh', 'Data awal dari data/warehouse_demo.db, inventory_dummy.xlsx, dan warehouse.db disalin. File sumber tidak diubah; file lama juga dapat berada di arsip/pengembangan_lama.'],
        ['Data yang digunakan aplikasi', 'Inventory sama dengan isi database uji warehouse_demo_seed.db. ID, tanggal, stok, dan lokasi barang yang sudah ada dipertahankan.'],
        ['Cara membuka data uji', 'Jalankan aplikasi paket uji yang menyertakan seed ini. Inventory sudah tersedia pada penggunaan pertama; tidak perlu mengimpor workbook yang sama.'],
        ['Cara impor ke database kosong', 'Daftarkan semua lokasi dari sheet Lokasi melalui Kelola Lokasi, kemudian pilih file ini di menu Impor / Ekspor, periksa pratinjau, dan simpan.'],
        ['Cegah duplikasi', 'Mengimpor workbook ini ke database yang sudah berisi ID sama akan ditolak. Gunakan Edit barang untuk memperbarui barang yang ada.'],
        ['Sheet yang dibaca aplikasi', 'Hanya Inventory dibaca saat impor barang. Sheet Lokasi, Petunjuk, sumber asli, dan riwayat merupakan referensi; riwayat tidak diimpor dari Excel.'],
        ['Kolom yang boleh diisi', 'Inventory kolom A:H berisi input. Jangan ubah header. Untuk barang baru, pakai Batch ID baru yang unik atau kosongkan A; Material tetap di B.'],
        ['Contoh baris baru', 'UJI-BARU-001 | MAT-NEW-001 | OIL FILTER ELEMENT | BHT | PBHT | A-B01-K01-R01-T01 | ' + today.isoformat() + ' | 10'],
        ['Format nilai', 'Kode memakai format Teks agar nol di depan tetap ada; GR Date YYYY-MM-DD; Total Stock bilangan bulat 0–999999999; seluruh kolom selain Batch ID wajib.'],
        ['Kode lokasi', 'A-B01-K03-R02-T02 berarti Area A, Rak B01, Kolom 03, Baris 02, Tingkat 02. Masing-masing rak terdiri dari 27 lokasi.'],
        ['Status umur', 'Aturan awal: Normal 0–30 hari; Warning 31–90 hari; Melewati batas mulai 91 hari. Status aplikasi dihitung kembali menurut tanggal komputer.'],
        ['Makna stok nol', 'Batch tetap tercatat, tetapi tidak menjadi stok aktif. D-D02 hanya berisi batch stok nol; D-D01 tidak memiliki batch sama sekali.'],
        ['Pemetaan Excel lama', 'Storage Bin asli dipertahankan di Excel_Lama dan Pemetaan_Bin. Kode lokasi barunya adalah penempatan SIMULASI, bukan hasil verifikasi lokasi fisik.'],
        ['Sumber data tambahan', 'Semua baris bertanda Simulasi tambahan pada Asal_Data adalah data rekaan untuk pengujian, bukan data perusahaan yang sudah diverifikasi.'],
        ['Skenario rekomendasi', 'Demand_Site_Simulasi, Kecocokan_Unit, dan Customer_Simulasi adalah bukti data rekaan untuk saran Site/Customer A/B/C. Data yang dipakai aplikasi tersimpan dalam database uji. Sheet ini hanya referensi, bukan jalur impor.'],
        ['Riwayat pemeriksaan', 'Sesi dan Hasil mempertahankan snapshot saat pengecekan. Status historis dapat berbeda dari status inventory hari ini.'],
        ['Riwayat aplikasi lama', 'Sesi_Lama dan Hasil_Lama berasal dari warehouse.db, disimpan apa adanya untuk referensi; tidak digabung menjadi sesi aktif aplikasi baru.'],
        ['Batas penggunaan', 'Workbook ini adalah data uji. Gunakan label dan hasilnya untuk demonstrasi aplikasi, bukan keputusan operasional atas stok nyata.'],
    ]
    g = _sheet(book, 'Petunjuk', ['Topik', 'Penjelasan'], guide)
    g.column_dimensions['A'].width = 31
    g.column_dimensions['B'].width = 120
    for row in range(2, g.max_row+1):
        g.row_dimensions[row].height = 42
    scenarios = [
        ['01', 'Batas umur 0, 29, 30 hari', 'UJI-AGE-000 / 029 / 030', 'Normal', 'Gunakan tanggal acuan di Petunjuk. Umur 30 masih Normal.'],
        ['02', 'Batas umur 31, 89, 90 hari', 'UJI-AGE-031 / 089 / 090', 'Warning', 'Umur 90 masih Warning; bukan Melewati batas.'],
        ['03', 'Batas umur 91 hari', 'UJI-AGE-091', 'Melewati batas', 'Status berubah setelah batas 90 hari.'],
        ['04', 'Rak hanya Normal', 'C-C02', 'Normal untuk stok aktif', 'Seluruh batch rak ini berumur 0–30 hari pada tanggal acuan.'],
        ['05', 'Rak hanya Warning', 'C-C03', 'Warning untuk stok aktif', 'Seluruh batch rak ini berumur 31–90 hari pada tanggal acuan.'],
        ['06', 'Rak hanya Melewati batas', 'C-C04', 'Melewati batas', 'Seluruh batch aktif lebih dari 90 hari pada tanggal acuan.'],
        ['07', 'Rak belum berisi barang', 'D-D01', 'Kosong', 'Lokasinya tersedia, tetapi tidak ada batch.'],
        ['08', 'Rak dengan stok habis', 'D-D02', 'Tanpa stok aktif', 'Batch tetap terlihat di Data Barang/filter stok habis.'],
        ['09', 'Tingkat berbeda dalam satu section', 'C-C01-K01-R01-T01 / T02 / T03', 'Normal / Warning / Melewati batas', 'UJI-MIX-T01..03: klik section untuk membandingkan ketiga tingkat.'],
        ['10', 'Beberapa batch di lokasi sama', 'UJI-SAME-001 / 002', 'Dua batch pada satu alamat', 'Lokasi sama, tanggal penerimaan berbeda, ID masing-masing tetap unik.'],
        ['11', 'Kode bernol depan', 'UJI-TEXT-001', 'Material 0000123405; Plant 0010', 'Nol depan bertahan saat dibaca, diekspor, dan diimpor.'],
        ['12', 'Deskripsi panjang', 'UJI-LONG-001', 'Teks lengkap dapat dibaca', 'Periksa detail barang, tabel, label, dan ekspor.'],
        ['13', 'Stok 0, 1, dan 15000', 'UJI-STOCK-000 / 001 / LARGE', 'Stok sesuai nilai input', 'Periksa filter stok dan pencatatan barang keluar.'],
        ['14', 'Riwayat tetap sama', 'Sesi / Hasil / Sesi_Lama / Hasil_Lama', 'Snapshot dan waktu lama tidak diubah', 'Pengecekan sebelumnya tidak dihitung ulang menjadi status hari ini.'],
        ['15', 'Pencegahan impor ganda', 'Inventory dengan ID yang sudah tersedia', 'Pratinjau menampilkan duplikasi', 'Tidak menambah ulang batch atau mengubah barang lama.'],
        ['16', 'Data Excel lama', 'LEGACY-001 sampai LEGACY-006; Pemetaan_Bin', 'Storage Bin asal tetap tersedia', 'Lokasi modern adalah penempatan simulasi yang dijelaskan eksplisit.'],
        ['17', 'Hijau: relokasi ke Site A', 'UJI-REC-001', 'Site A', 'Demand naik dan open demand melebihi stok simulasi Site A.'],
        ['18', 'Hijau: belum perlu tindakan', 'UJI-REC-002', 'Tidak ada tujuan', 'Kebutuhan Site A/B/C tidak melampaui stoknya.'],
        ['19', 'Kuning: tawarkan ke Customer A', 'UJI-REC-003', 'Customer A', 'Tetap memerlukan cek fisik sebelum penawaran.'],
        ['20', 'Merah: tawarkan ke Customer A', 'UJI-REC-004', 'Customer A', 'Periksa kelayakan fisik terlebih dahulu.'],
        ['21', 'Merah: redistribusi Site B', 'UJI-REC-005', 'Site B', 'Demand Site B tinggi; periksa fisik sebelum transfer.'],
        ['22', 'Merah: dua pilihan', 'UJI-REC-006', 'Customer A / Site C', 'Dua alternatif dapat dibandingkan petugas sesudah cek fisik.'],
    ]
    sc = _sheet(book, 'Skenario_Uji', ['No.', 'Skenario', 'Data / lokasi', 'Hasil yang diharapkan', 'Catatan'], scenarios)
    for col, width in {'A': 8, 'B': 33, 'C': 46, 'D': 36, 'E': 80}.items():
        sc.column_dimensions[col].width = width
    for row in range(2, sc.max_row+1):
        sc.row_dimensions[row].height = 45
    for name, table in [('Demand_Site_Simulasi','site_demand'),
                        ('Kecocokan_Unit','material_unit_compatibility'),
                        ('Customer_Simulasi','customer_demand')]:
        _dict_sheet(book, name, table_rows(db, table),
                    [r['name'] for r in db.execute(f'PRAGMA table_info("{table}")')])
    origin = _sheet(book, 'Asal_Data', ['Batch ID', 'Jenis sumber', 'File sumber', 'Baris sumber',
                                     'Storage Bin asli', 'Location Code', 'Catatan'], provenance)
    origin.column_dimensions['A'].width = 43
    origin.column_dimensions['F'].width = 30
    origin.column_dimensions['G'].width = 80
    _sheet(book, 'Pemetaan_Bin', ['Baris Excel', 'Material', 'Storage Bin asli', 'Batch ID',
                                'Location Code simulasi', 'Penjelasan'], mapping)
    _dict_sheet(book, 'Barang_Asli', originals['batches'], FIELDS)
    _sheet(book, 'Excel_Lama', ['Baris Excel'] + legacy_headers,
           ([number] + [row.get(key) for key in legacy_headers] for number, row in legacy_records))
    source_names = {
        'check_sessions': 'Sesi', 'check_results': 'Hasil',
        'stock_movements': 'Barang_Keluar', 'batch_location_migration': 'Migrasi_Lokasi',
        'metadata': 'Metadata_Asli',
    }
    for table, name in source_names.items():
        headers = [r['name'] for r in db.execute(f'PRAGMA table_info("{table}")')]
        ws = _dict_sheet(book, name, originals[table], headers)
        for i, cell in enumerate(ws[1], 1):
            if cell.value in ('snapshot', 'note', 'value'):
                ws.column_dimensions[get_column_letter(i)].width = 95
            elif cell.value in ('batch_id', 'id', 'old_location_code', 'new_location_code'):
                ws.column_dimensions[get_column_letter(i)].width = 42
    for table, name in [('sessions', 'Sesi_Lama'), ('scans', 'Hasil_Lama')]:
        ws = _dict_sheet(book, name, old_history[table]['rows'], old_history[table]['headers'])
        for i, cell in enumerate(ws[1], 1):
            if cell.value in ('snapshot', 'source', 'item_key'):
                ws.column_dimensions[get_column_letter(i)].width = 100
    book.active = 0
    book.save(path)
    book.close()


def validate_workbook(path, seed_db):
    records = read_inventory(path)
    with tempfile.TemporaryDirectory(prefix='warehouse-trial-validate-') as folder:
        fresh = connect(Path(folder)/'fresh.db')
        try:
            with fresh:
                fresh.executemany('''INSERT INTO locations(code,area,rack,"column","row",level)
                    VALUES(?,?,?,?,?,?)''', [tuple(r[k] for k in LOCATION_FIELDS)
                                           for r in seed_db.execute('SELECT * FROM locations')])
            errors = [row for row in preview(fresh, records) if row['errors']]
            if errors:
                raise ValueError('Workbook tidak lolos validasi impor: ' + repr(errors[:3]))
            imported = commit_import(fresh, records)
            actual = {r['id']: tuple(r[k] for k in FIELDS) for r in fresh.execute('SELECT * FROM batches')}
            expected = {r['id']: tuple(r[k] for k in FIELDS) for r in seed_db.execute('SELECT * FROM batches')}
            if actual != expected:
                raise ValueError('Isi workbook berbeda dari database uji.')
            duplicate_rows = preview(fresh, records)
            if not all(any('duplikat' in error for error in row['errors']) for row in duplicate_rows):
                raise ValueError('Pemeriksaan duplikasi tidak sesuai harapan.')
            return {'imported_batches': len(imported), 'duplicate_ids_rejected': len(duplicate_rows)}
        finally:
            fresh.close()


def build_trial_dataset(root=PROJECT_ROOT, output_dir=None, today=None, target_count=200):
    root = Path(root).resolve()
    output_dir = Path(output_dir or root/'data_uji').resolve()
    today = today or date.today()
    if today > date.today():
        raise ValueError('Tanggal acuan tidak boleh setelah hari ini: impor aplikasi menolak tanggal masa depan.')
    source_path = root/'data'/'warehouse_demo.db'
    def old_source(name):
        current = root/name
        return current if current.is_file() else root/'arsip'/'pengembangan_lama'/name

    legacy_path = old_source('inventory_dummy.xlsx')
    legacy_db_path = old_source('warehouse.db')
    for source in (source_path, legacy_path, legacy_db_path):
        if not source.is_file():
            raise FileNotFoundError(f'Sumber tidak ditemukan: {source}')
    hashes_before = {str(path.relative_to(root)): file_hash(path)
                     for path in (source_path, legacy_path, legacy_db_path)}
    output_dir.mkdir(parents=True, exist_ok=True)
    seed_path = output_dir/'warehouse_demo_seed.db'
    if seed_path in (source_path, legacy_db_path):
        raise ValueError('Output harus berbeda dari database sumber.')
    legacy_headers, legacy_records = read_legacy_inventory(legacy_path)
    old = readonly_database(legacy_db_path)
    try:
        old_history = {table: {'headers': [r['name'] for r in old.execute(f'PRAGMA table_info("{table}")')],
                              'rows': table_rows(old, table)} for table in ('sessions', 'scans')}
    finally:
        old.close()
    # Build next to the destination and replace only after all checks pass.
    with tempfile.TemporaryDirectory(prefix='.build-', dir=output_dir) as folder:
        working_seed = Path(folder)/'warehouse_demo_seed.db'
        source_db = readonly_database(source_path)
        db = sqlite3.connect(working_seed)
        db.row_factory = sqlite3.Row
        try:
            source_db.backup(db)
        finally:
            source_db.close()
            db.close()
        # Existing source files may still use v5. Upgrade only the disposable
        # copy so the user's source database and history stay untouched.
        db = connect(working_seed)
        try:
            db.execute('PRAGMA foreign_keys=ON')
            tables = ('batches', 'locations', 'check_sessions', 'check_results', 'stock_movements',
                      'batch_location_migration', 'metadata', 'site_demand',
                      'material_unit_compatibility', 'customer_demand')
            originals = {table: table_rows(db, table) for table in tables}
            if len(originals['batches']) >= target_count:
                raise ValueError('Target harus lebih besar dari jumlah batch sumber agar skenario tambahan tersedia.')
            if any(date.fromisoformat(row['received_date']) > today for row in originals['batches']):
                raise ValueError('Tanggal acuan mendahului tanggal penerimaan sumber; pilih tanggal acuan yang lebih baru.')
            provenance = [[r['id'], 'Data tersimpan', 'data/warehouse_demo.db', '', '', r['location_code'],
                           'Seluruh nilai asli dipertahankan.'] for r in originals['batches']]
            mapping = []
            with db:
                for area, rack in ADDED_RACKS:
                    for column in range(1, 4):
                        for row in range(1, 4):
                            for level in range(1, 4):
                                code = location_code(area, rack, column, row, level)
                                db.execute('''INSERT OR IGNORE INTO locations(code,area,rack,"column","row",level)
                                    VALUES(?,?,?,?,?,?)''', (code, area, rack, column, row, level))
                for i, (line, record) in enumerate(legacy_records, 1):
                    values = dict(material=str(record['Material']), description=str(record['Material Description']),
                                  plant=str(record['Plant']), storage_location=str(record['Storage Location']),
                                  received_date=_date_text(record['GR Date']), stock=int(record['Total Stock']))
                    matches = [r for r in originals['batches']
                               if all(r[key] == value for key, value in values.items())]
                    if len(matches) == 1:
                        key, code = matches[0]['id'], matches[0]['location_code']
                        detail = 'Nilai barang cocok dengan satu batch tersimpan; tidak diduplikasi. Storage Bin asal belum diverifikasi secara fisik.'
                    else:
                        key = f'LEGACY-{i:03d}'
                        code = location_code('C', 'C01', (i-1)//9+1, (i-1)%3+1, 3)
                        _insert_batch(db, dict(values, id=key, location_code=code))
                        detail = 'Lokasi penempatan SIMULASI; Storage Bin asli tidak menyatakan kolom/baris/tingkat.'
                        provenance.append([key, 'Excel lama', 'inventory_dummy.xlsx', line,
                                           str(record['Storage Bin']), code, detail])
                    mapping.append([line, record['Material'], str(record['Storage Bin']), key, code, detail])
                for age in (0, 29, 30, 31, 89, 90, 91):
                    rack = 'C02' if age <= 30 else 'C03' if age <= 90 else 'C04'
                    _add_synthetic(db, f'UJI-AGE-{age:03d}', 'AGE-'+str(age).zfill(3),
                                   f'CONTOH BATAS UMUR {age} HARI', location_code('C', rack, 1, 1, age%3+1),
                                   age, 1, today, provenance, f'Batas umur {age} hari pada {today.isoformat()}.')
                for level, age in [(1, 10), (2, 50), (3, 120)]:
                    _add_synthetic(db, f'UJI-MIX-T{level:02d}', 'MIX-'+str(level), 'CONTOH TINGKAT BERBEDA',
                                   location_code('C','C01',1,1,level), age, 5, today, provenance,
                                   'Satu section; tingkat bawah Normal, tengah Warning, atas Melewati batas.')
                for i, age in [(1, 15), (2, 110)]:
                    _add_synthetic(db, f'UJI-SAME-{i:03d}', 'HYD-SAME-001', 'HOSE ASSY - PENERIMAAN BERBEDA',
                                   location_code('C','C01',2,1,1), age, 12, today, provenance,
                                   'Material dan lokasi sama; ID batch dan tanggal penerimaan berbeda.')
                _add_synthetic(db, 'UJI-TEXT-001', '0000123405', 'SEAL KIT - KODE BERNOL DEPAN',
                               location_code('C','C02',3,3,1), 12, 24, today, provenance,
                               'Kode material, Plant, dan Storage Location harus tetap Teks.', plant='0010', storage='0001')
                _add_synthetic(db, 'UJI-LONG-001', 'HYD-LONG-001',
                               'HYDRAULIC HOSE ASSEMBLY HIGH PRESSURE 1/2 INCH - KONEKTOR SIKU 90 DERAJAT, '
                               'PANJANG 1500 MM, UNTUK SIMULASI PEMERIKSAAN KETERBACAAN DESKRIPSI BARANG DI TABEL DAN LABEL',
                               location_code('C','C03',3,3,1), 55, 6, today, provenance, 'Deskripsi panjang, tetap di bawah 200 karakter.')
                for key, stock, age in [('000', 0, 120), ('001', 1, 12), ('LARGE', 15000, 60)]:
                    rack, area = ('D02','D') if stock == 0 else ('C01','C')
                    _add_synthetic(db, 'UJI-STOCK-'+key, 'STOCK-'+key, 'CONTOH JUMLAH STOK '+str(stock),
                                   location_code(area,rack,3,2,1), age, stock, today, provenance, 'Uji stok nol, satu, atau besar.')
                for i in range(1, 10):
                    _add_synthetic(db, f'UJI-ZERO-{i:03d}', f'ZERO-{i:03d}', 'STOK HABIS - DATA TETAP TERCATAT',
                                   location_code('D','D02',(i-1)//3+1,(i-1)%3+1,1),
                                   [10,60,150][(i-1)%3], 0, today, provenance, 'Batch stok nol tidak menjadi stok aktif.')
                for index, (age, stock, label) in enumerate([
                    (12,20,'Normal: demand Site A meningkat, kandidat relokasi.'),
                    (12,8,'Normal: kebutuhan site cukup, tidak perlu tindakan.'),
                    (60,7,'Warning: cek fisik, tawarkan ke Customer A.'),
                    (120,10,'Melewati batas: cek fisik, tawarkan ke Customer A.'),
                    (180,15,'Melewati batas: cek fisik, redistribusi ke Site B.'),
                    (240,20,'Melewati batas: pilihan Customer A dan Site C.'),
                ], 1):
                    _add_synthetic(db, f'UJI-REC-{index:03d}', f'MAT-REC-{index:03d}',
                                   f'CONTOH REKOMENDASI {index:02d}',
                                   location_code('A','B01',(index-1)//3+1,(index-1)%3+1,1),
                                   age, stock, today, provenance, label)
                # Fill each status-only rack with varied rows/columns/levels.
                for rack, ages in [('C02',[0,7,15,29,30]), ('C03',[31,45,60,89,90]), ('C04',[91,120,180,365,730])]:
                    for i in range(1, 16):
                        offset = (i*10)%27
                        code = location_code('C',rack,offset//9+1,(offset%9)//3+1,offset%3+1)
                        part, description = PARTS[(i-1)%len(PARTS)]
                        _add_synthetic(db, f'UJI-{rack}-{i:03d}', part, description, code,
                                       ages[(i-1)%len(ages)], [1,2,5,12,50][(i-1)%5], today,
                                       provenance, 'Rak khusus satu status pada tanggal acuan.')
                remaining = target_count - db.execute('SELECT COUNT(*) FROM batches').fetchone()[0]
                if remaining < 0:
                    raise ValueError('Target terlalu kecil untuk mencakup seluruh skenario tambahan.')
                for i in range(1, remaining+1):
                    area, rack = EXTRA_RACKS[(i-1)%len(EXTRA_RACKS)]
                    offset = (i*10)%27
                    part, description = PARTS[(i-1)%len(PARTS)]
                    _add_synthetic(db, f'UJI-INV-{i:03d}', part, description,
                                   location_code(area,rack,offset//9+1,(offset%9)//3+1,offset%3+1),
                                   [2,18,29,31,45,65,89,90,91,120,240,365,730][(i-1)%13],
                                   [1,2,4,8,16,25,50,100,250][(i-1)%9], today, provenance,
                                   'Sebaran inventory simulasi: beberapa material, penerimaan, dan lokasi.')
                _seed_recommendation_scenarios(db, today)
                db.execute('INSERT OR REPLACE INTO metadata(key,value) VALUES(?,?)', ('trial_data_as_of',today.isoformat()))
                db.execute('INSERT OR REPLACE INTO metadata(key,value) VALUES(?,?)', ('trial_data_version','1'))
            for table in tables:
                actual = table_rows(db, table)
                if any(row not in actual for row in originals[table]):
                    raise ValueError('Data sumber berubah pada salinan: '+table)
            if db.execute('PRAGMA integrity_check').fetchone()[0] != 'ok' or db.execute('PRAGMA foreign_key_check').fetchall():
                raise ValueError('Integritas database uji gagal.')
            working_xlsx = Path(folder)/'Data_Uji_Warehouse.xlsx'
            _workbook(working_xlsx, db, originals, legacy_headers, legacy_records,
                      old_history, mapping, provenance, today)
            validation = validate_workbook(working_xlsx, db)
            batches = table_rows(db, 'batches')
            counts = Counter(aging(row['received_date'], today=today)[1] for row in batches if row['stock']>0)
            manifest = {
                'reference_date': today.isoformat(), 'seed_schema': db.execute('PRAGMA user_version').fetchone()[0],
                'inventory_batches': len(batches), 'original_batches_preserved': len(originals['batches']),
                'legacy_excel_rows': len(legacy_records),
                'legacy_batches_added': sum(r[1]=='Excel lama' for r in provenance),
                'synthetic_batches_added': sum(r[1]=='Simulasi tambahan' for r in provenance),
                'locations': db.execute('SELECT COUNT(*) FROM locations').fetchone()[0],
                'racks': db.execute('SELECT COUNT(*) FROM (SELECT DISTINCT area,rack FROM locations)').fetchone()[0],
                'active_batches': sum(row['stock']>0 for row in batches),
                'zero_stock_batches': sum(row['stock']==0 for row in batches),
                'status_active_at_reference_date': dict(counts),
                'current_sessions_preserved': len(originals['check_sessions']),
                'current_results_preserved': len(originals['check_results']),
                'legacy_sessions_in_workbook': len(old_history['sessions']['rows']),
                'legacy_results_in_workbook': len(old_history['scans']['rows']),
                'validation': validation, 'source_sha256': hashes_before,
            }
        finally:
            db.close()
        hashes_after = {str(path.relative_to(root)): file_hash(path)
                        for path in (source_path, legacy_path, legacy_db_path)}
        if hashes_before != hashes_after:
            raise RuntimeError('File sumber berubah saat generator berjalan. Output belum dipublikasikan; ulangi setelah aplikasi ditutup.')
        working_seed.replace(seed_path)
        working_xlsx.replace(output_dir/'Data_Uji_Warehouse.xlsx')
        manifest['output_sha256'] = {name: file_hash(output_dir/name)
                                     for name in ('warehouse_demo_seed.db','Data_Uji_Warehouse.xlsx')}
        (output_dir/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--today', type=date.fromisoformat, default=date.today(), help='Tanggal acuan YYYY-MM-DD')
    parser.add_argument('--output-dir', type=Path, default=PROJECT_ROOT/'data_uji')
    args = parser.parse_args()
    print(json.dumps(build_trial_dataset(output_dir=args.output_dir, today=args.today), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
