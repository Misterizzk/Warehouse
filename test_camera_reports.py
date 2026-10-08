import tempfile
import unittest
import threading
from pathlib import Path
from datetime import date
from unittest.mock import patch
import cv2
import numpy as np
import pypdfium2 as pdfium
from openpyxl import load_workbook
from warehouse.database import connect,seed_demo
from warehouse.sessions import Sessions
from warehouse.inventory import Inventory
from warehouse.labels import qr_image
from warehouse.camera import decode_frame,open_device,CameraWorker,DEVICE_LOCK
from warehouse.reports import history,report_data,export_excel,export_pdf


def qr_frame(key='DEMO-0001'):
    image=np.full((720,1280,3),245,dtype=np.uint8)
    qr=np.array(qr_image(key))
    image[100:100+qr.shape[0],100:100+qr.shape[1]]=qr
    return image


class CameraTests(unittest.TestCase):
    def test_decode_real_qr(self):
        found=decode_frame(qr_frame())
        self.assertEqual(found[0][0],'ID:DEMO-0001')
        self.assertEqual(found[0][1].shape,(4,2))
        self.assertEqual(decode_frame(np.zeros((480,640,3),dtype=np.uint8)),[])

    def test_decode_multiple_real_qrs_in_one_frame(self):
        frame=np.full((720,1280,3),245,dtype=np.uint8)
        for key,x in (('DEMO-0001',80),('DEMO-0002',760)):
            qr=np.array(qr_image(key))
            frame[120:120+qr.shape[0],x:x+qr.shape[1]]=qr
        self.assertEqual({code for code,_ in decode_frame(frame)},
                         {'ID:DEMO-0001','ID:DEMO-0002'})

    def test_failed_capture_release(self):
        class Failed:
            released=0
            def __init__(self,*args):pass
            def isOpened(self):return False
            def release(self):Failed.released+=1
        self.assertIsNone(open_device(0,factory=Failed))
        self.assertGreater(Failed.released,0)
        self.assertIsNone(open_device(0,stop=self.stopped(),factory=Failed))

    def stopped(self):
        event=threading.Event()
        event.set()
        return event

    def test_worker_disconnect_and_cleanup(self):
        class Capture:
            released=False
            def read(self):return False,None
            def release(self):self.released=True
        cap=Capture()
        worker=CameraWorker()
        with patch('warehouse.camera.open_device',return_value=(cap,qr_frame())):
            worker.start()
            worker.thread.join(5)
        self.assertFalse(worker.thread.is_alive())
        self.assertTrue(cap.released)
        self.assertFalse(DEVICE_LOCK.locked())
        messages=[]
        while not worker.messages.empty():messages.append(worker.messages.get()[0])
        self.assertIn('error',messages)
        self.assertIn('stopped',messages)


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.db=connect(Path(self.temp.name)/'demo.db')
        seed_demo(self.db)
        self.service=Sessions(self.db)
        self.key=self.service.start('Faiz','Area A')

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def pdf_text(self,path):
        pdf=pdfium.PdfDocument(path)
        texts=[]
        count=len(pdf)
        for i in range(count):
            page=pdf[i]
            text=page.get_textpage()
            texts.append(text.get_text_range())
            text.close()
            page.close()
        pdf.close()
        return '\n'.join(texts),count

    def test_filters_and_empty_draft_pdf(self):
        day=date.today().isoformat()
        self.assertEqual(len(history(self.db,query='faiz',status='Draft',date_from=day,date_to=day)),1)
        self.assertEqual(history(self.db,status='Selesai'),[])
        for values in ({'date_from':'bad'},{'date_from':'2027-01-01','date_to':'2026-01-01'}):
            with self.assertRaises(ValueError):history(self.db,**values)
        path=Path(self.temp.name)/'empty.pdf'
        export_pdf(report_data(self.db,self.key),path)
        text,_=self.pdf_text(path)
        self.assertIn('DRAFT',text)
        self.assertIn('Belum ada hasil',text)

    def test_snapshot_exports_full_notes_and_escaping(self):
        for code in ('DEMO-0001','DEMO-0026','DEMO-0041'):self.service.add_code(self.key,code)
        note='=2+2 <periksa & kemasan>\n'+'Catatan lengkap. '*100+'AKHIR CATATAN'
        self.service.save_note(self.key,note,'DEMO-0001')
        self.service.save_note(self.key,'Catatan sesi <aman> & tetap')
        self.service.finish(self.key)
        before=report_data(self.db,self.key)
        item=Inventory(self.db).list_batches(query='DEMO-0001')[0]
        Inventory(self.db).update(item['id'],dict(item,material='CHANGED',description='Changed',stock='999'))
        data=report_data(self.db,self.key)
        self.assertEqual(before['rows'],data['rows'])
        xlsx=Path(self.temp.name)/'report.xlsx'
        export_excel(data,xlsx)
        book=load_workbook(xlsx)
        self.assertEqual(book['Hasil pengecekan']['M2'].value,note)
        self.assertEqual(book['Hasil pengecekan']['M2'].data_type,'s')
        self.assertEqual(book['Hasil pengecekan']['D2'].value,item['material'])
        self.assertEqual(book['Ringkasan']['B3'].value,'Selesai')
        book.close()
        pdf=Path(self.temp.name)/'report.pdf'
        export_pdf(data,pdf)
        text,_=self.pdf_text(pdf)
        for expected in ('CHK-000001','Faiz','Melewati batas','AKHIR CATATAN','<periksa & kemasan>'):
            self.assertIn(expected,text)
        self.assertNotIn('CHANGED',text)

    def test_multipage_headers_and_notes(self):
        for row in Inventory(self.db).list_batches():self.service.add_code(self.key,row['id'])
        self.service.save_note(self.key,'A'*2000,'DEMO-0050')
        path=Path(self.temp.name)/'long.pdf'
        export_pdf(report_data(self.db,self.key),path)
        text,count=self.pdf_text(path)
        self.assertGreater(count,1)
        self.assertIn('DEMO-0050',text)
        self.assertEqual(text.count('Halaman '),count)


if __name__=='__main__':unittest.main()
