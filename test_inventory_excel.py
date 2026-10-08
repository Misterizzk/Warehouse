import tempfile
import unittest
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from unittest.mock import patch
from warehouse.database import connect,seed_demo
from warehouse.inventory_excel import write_inventory,read_inventory,preview,commit_import


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.path=Path(self.temp.name)/'inventory.xlsx'
        self.db=connect(Path(self.temp.name)/'test.db')
        seed_demo(self.db)

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def input_file(self,rows):
        write_inventory(self.path,self.db,True)
        book=load_workbook(self.path)
        sheet=book['Inventory']
        sheet.delete_rows(2,sheet.max_row)
        for row in rows:sheet.append(row)
        book.save(self.path)
        book.close()

    def row(self,key='',material='0007'):
        return [key,material,'New part','BHT','PBHT','A-B01-01-01',date.today(),'3']

    def test_template_import_atomic_and_repeated(self):
        self.input_file([self.row(),self.row('EXT-002','0008')])
        source=read_inventory(self.path)
        checked=preview(self.db,source)
        self.assertFalse(any(r['errors'] for r in checked))
        self.assertEqual(checked[0]['values']['location_code'],'A-B01-K01-R01-T01')
        self.assertTrue(checked[0]['notices'])
        keys=commit_import(self.db,source)
        self.assertEqual(len(keys),2)
        self.assertEqual(self.db.execute('SELECT material FROM batches WHERE id=?',(keys[0],)).fetchone()[0],'0007')
        with self.assertRaises(ValueError):commit_import(self.db,source)
        self.assertEqual(self.db.execute('SELECT count(*) FROM batches').fetchone()[0],52)

    def test_errors_and_duplicate_no_partial_write(self):
        invalid=self.row('DUP','0008')
        invalid[5]='missing'
        self.input_file([self.row('DUP'),invalid,self.row('DEMO-0001')])
        source=read_inventory(self.path)
        self.assertTrue(preview(self.db,source)[1]['errors'])
        with self.assertRaises(ValueError):commit_import(self.db,source)
        self.assertEqual(self.db.execute('SELECT count(*) FROM batches').fetchone()[0],50)

    def test_formula_numeric_codes_and_identical_rows(self):
        numeric=self.row('N')
        numeric[1]=7
        formula=self.row('F')
        formula[7]='=1+2'
        self.input_file([numeric,formula,self.row(),self.row()])
        checked=preview(self.db,read_inventory(self.path))
        for i in (0,1,3):self.assertTrue(checked[i]['errors'])

    def test_export_safe_text_and_roundtrip_ids(self):
        with self.db:self.db.execute("UPDATE batches SET description='=2+2' WHERE id='DEMO-0001'")
        write_inventory(self.path,self.db)
        book=load_workbook(self.path)
        self.assertEqual(book['Inventory']['C2'].data_type,'s')
        self.assertEqual(book['Inventory'].max_row,51)
        self.assertEqual([cell.value for cell in book['Lokasi'][1]],
                         ['Location Code','Area','Rak','Kolom','Baris','Tingkat'])
        book.close()
        checked=preview(self.db,read_inventory(self.path))
        self.assertEqual(len(checked),50)
        self.assertTrue(all(r['errors'] for r in checked))

    def test_commit_revalidates_changes_after_preview(self):
        self.input_file([self.row('RACE')])
        source=read_inventory(self.path)
        self.assertFalse(preview(self.db,source)[0]['errors'])
        commit_import(self.db,source)
        with self.assertRaises(ValueError):commit_import(self.db,source)

    def test_unknown_sheet_dimensions_and_shifted_first_row(self):
        shifted=['MAT-ELC-001','Kabel USB Type-C 1 Meter','A','B01','A-B01-01-01',
                 '2026-08-15',250]
        valid=['','MAT-ELC-002','Mouse Wireless','A','B01','A-B01-01-02',
               date.today(),120]
        self.input_file([shifted,valid])
        original=load_workbook
        def unknown_dimensions(*args,**kwargs):
            book=original(*args,**kwargs)
            book['Inventory']._max_row=None
            return book
        with patch('warehouse.inventory_excel.load_workbook',side_effect=unknown_dimensions):
            source=read_inventory(self.path)
        self.assertEqual(len(source),2)
        checked=preview(self.db,source)
        self.assertIn('Kolom tampak bergeser',' '.join(checked[0]['errors']))
        self.assertFalse(checked[1]['errors'],checked[1]['errors'])
        with self.assertRaises(ValueError):commit_import(self.db,source)
        self.assertEqual(self.db.execute('SELECT count(*) FROM batches').fetchone()[0],50)


if __name__=='__main__':unittest.main()
