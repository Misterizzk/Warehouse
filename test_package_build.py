"""A new Windows package starts with demo inventory but no prior officer sessions."""
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from openpyxl import load_workbook

from scripts.build_windows import create_clean_install_database, create_clean_install_workbook
from warehouse.database import SCHEMA_VERSION


ROOT = Path(__file__).resolve().parent


class CleanPackageDatabaseTests(unittest.TestCase):
    def test_install_copy_keeps_inventory_and_removes_old_checks(self):
        source = ROOT / "data_uji" / "warehouse_demo_seed.db"
        with closing(sqlite3.connect(source.resolve().as_uri() + "?mode=ro", uri=True)) as db:
            expected = db.execute("SELECT COUNT(*) FROM batches").fetchone()[0]
            old_sessions = db.execute("SELECT COUNT(*) FROM check_sessions").fetchone()[0]
        self.assertGreater(old_sessions, 0)
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "warehouse_demo.db"
            create_clean_install_database(source, target)
            workbook_path = Path(folder) / "Data_Uji_Warehouse.xlsx"
            create_clean_install_workbook(target, workbook_path)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, [
                    "Inventory", "Lokasi", "Petunjuk", "Demand_Site_Simulasi",
                    "Kecocokan_Unit", "Customer_Simulasi"])
                self.assertEqual(workbook["Inventory"].max_row - 1, expected)
                self.assertIn(("MAT-REC-001", "A", 3, 9, 10, 1),
                              list(workbook["Demand_Site_Simulasi"].values))
                self.assertIn(("MAT-REC-003", "UNIT-CUSTOMER-A-WARNING"),
                              list(workbook["Kecocokan_Unit"].values))
                self.assertIn(("A", "UNIT-CUSTOMER-A-WARNING", 4),
                              list(workbook["Customer_Simulasi"].values))
            finally:
                workbook.close()
            with closing(sqlite3.connect(target)) as db:
                self.assertEqual(db.execute("PRAGMA user_version").fetchone()[0], SCHEMA_VERSION)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM batches").fetchone()[0], expected)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM check_sessions").fetchone()[0], 0)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM check_results").fetchone()[0], 0)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM recommendation_decisions").fetchone()[0], 0)
                self.assertIsNotNone(db.execute("SELECT value FROM metadata WHERE key='seed_date'").fetchone())
                self.assertEqual(db.execute("PRAGMA foreign_key_check").fetchall(), [])
        with closing(sqlite3.connect(source.resolve().as_uri() + "?mode=ro", uri=True)) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM check_sessions").fetchone()[0], old_sessions)


if __name__ == "__main__":
    unittest.main()
